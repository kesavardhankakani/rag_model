from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import List

import fitz
from docx import Document


@dataclass
class DocumentPage:
    number: int
    text: str


@dataclass
class ExtractedDocument:
    filename: str
    file_type: str
    pages: List[DocumentPage] = field(default_factory=list)
    text: str = ""
    metadata: dict = field(default_factory=dict)

    @property
    def page_count(self) -> int:
        return len(self.pages)

    def to_dict(self):
        return {
            "filename": self.filename,
            "file_type": self.file_type,
            "pages": [{"number": page.number, "text": page.text} for page in self.pages],
            "text": self.text,
        }


def _extract_pdf(path: str | Path) -> ExtractedDocument:
    doc = fitz.open(path)
    pages: List[DocumentPage] = []
    all_text: List[str] = []

    for page_index in range(len(doc)):
        page = doc[page_index]
        page_text = page.get_text("text") or ""
        pages.append(DocumentPage(number=page_index + 1, text=page_text))
        all_text.append(page_text)

    doc.close()
    combined_text = "\n\n".join(all_text).strip()
    return ExtractedDocument(
        filename=Path(path).name,
        file_type="pdf",
        pages=pages,
        text=combined_text,
        metadata={"page_count": len(pages)},
    )


def _extract_docx(path: str | Path) -> ExtractedDocument:
    document = Document(str(path))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text and paragraph.text.strip()]
    text = "\n\n".join(paragraphs)
    return ExtractedDocument(
        filename=Path(path).name,
        file_type="docx",
        pages=[DocumentPage(number=1, text=text)] if text else [],
        text=text,
        metadata={"page_count": 1},
    )


def extract_document(file_path: str | Path) -> ExtractedDocument:
    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)
    raise ValueError(f"Unsupported file type: {suffix}. Upload a PDF or DOCX file.")
